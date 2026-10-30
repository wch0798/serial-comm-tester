"""
시험 보고서 (엑셀 .xlsx)

시나리오 실행 결과를 엑셀로 남긴다. 화면에서도, 화면 없는 자동 실행(autotest.py)에서도 같이 쓴다.
시트 구성
- 요약      : 전체 판정(합격/불합격), 연결별 통과·실패 수
- 실패 목록 : 실패한 단계만 모아 보기
- 단계 결과 : 모든 단계의 결과 (회차, 걸린 시간, 메시지)
- 통신 기록 : 시험 중 주고받은 데이터 (시각, 구분, HEX, 텍스트)
"""

import datetime
import os
import platform
import time

import serial_core as core
import scenario as sc

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None

MAX_COMM = 50000  # 통신 기록 최대 행 수 (넘으면 앞부분을 버린다)

KIND_NAMES = {"rx": "받음", "tx": "보냄", "auto": "자동응답", "relay": "중계", "fault": "오류주입",
              "error": "오류", "info": "정보", "decode": "해석", "scenario": "시나리오"}

# 로그 종류별 색 (글자/배지색, 줄 배경색) - 화면 로그와 보고서가 같이 쓴다.
# 종류마다 색상(색깔 방향)을 크게 떨어뜨려 겹치지 않게 골랐다:
# 수신 파랑 · 송신 초록 · 자동응답 주황갈색 · 중계 청록 · 해석 보라 · 오류주입 자주(분홍보라)
# · 시나리오 황토 · 정보 회색 · 오류 빨강.  배지는 흰 글자, 줄 배경 위 글자 모두 읽기 쉬운 대비(4.5 이상)
KIND_COLORS = {
    "rx": ("#0F60D1", "#DDEBFD"),
    "tx": ("#227514", "#E2F7DD"),
    "auto": ("#9F510E", "#FEEBD8"),
    "relay": ("#03726F", "#D9F7F5"),
    "decode": ("#7A2FE0", "#EFE4FF"),
    "fault": ("#B81A98", "#FBDFF5"),
    "scenario": ("#7B620F", "#F8F1DA"),
    "info": ("#5F6368", "#E3E5E8"),
    "error": ("#CB1B34", "#FFDDE2"),
}

PASS_TEXT = "합격"
FAIL_TEXT = "불합격"

_GREEN = "DCFCE7"
_RED = "FEE2E2"
_HEAD = "1F3A5F"
_LIGHT = "EEF2F7"


class ReportError(Exception):
    pass


def ts(t, ms=True):
    if t is None:
        return ""
    d = datetime.datetime.fromtimestamp(t)
    return d.strftime("%Y-%m-%d %H:%M:%S.%f")[:-3] if ms else d.strftime("%Y-%m-%d %H:%M:%S")


def duration_text(sec):
    if sec is None:
        return ""
    sec = max(0.0, sec)
    if sec < 60:
        return "%.2f초" % sec
    m, s = divmod(int(sec), 60)
    h, m = divmod(m, 60)
    return ("%d시간 %d분 %d초" % (h, m, s)) if h else ("%d분 %d초" % (m, s))


class CommLog(object):
    """시험 중 주고받은 데이터를 모은다 (여러 스레드에서 add 해도 된다)."""

    def __init__(self, limit=MAX_COMM):
        self.rows = []
        self.limit = limit
        self.dropped = 0

    def add(self, tab_name, kind, data, src=None, now=None):
        if kind not in KIND_NAMES:
            return
        if len(self.rows) >= self.limit:
            cut = max(1, self.limit // 10)
            del self.rows[:cut]
            self.dropped += cut
        self.rows.append((now or time.time(), tab_name, kind, data, src))


def tab_result(name, conn_text, runner, repeat=None):
    """ScenarioRunner → 보고서용 dict"""
    end = runner.ended_at
    if end is None and runner.records:
        end = runner.records[-1]["time"]
    return dict(
        name=name, conn=conn_text, started=runner.started_at, ended=end,
        end_reason=runner.end_reason or ("실행 중" if runner.running else ""),
        loops=runner.loop, repeat=runner.repeat if repeat is None else repeat,
        passed=runner.passed, failed=runner.failed, ok=runner.ok(),
        records=list(runner.records), dropped=runner.dropped)


def error_result(name, conn_text, message):
    """연결을 못 열었거나 시작조차 못 한 경우"""
    now = time.time()
    return dict(name=name, conn=conn_text, started=now, ended=now, end_reason=message, loops=0, repeat=0,
                passed=0, failed=1, ok=False, records=[], dropped=0)


def overall_ok(results):
    return bool(results) and all(r["ok"] for r in results)


def _data_cells(kind, data):
    """통신 기록 한 줄의 (HEX, 텍스트)"""
    if isinstance(data, (bytes, bytearray)):
        data = bytes(data)
        if kind in ("info", "error", "fault", "decode", "scenario"):
            return "", data.decode("utf-8", "replace")
        return core.bytes_to_hex(data), core.bytes_to_text(data)
    return "", "" if data is None else str(data)


def _src_text(src):
    if src is None:
        return ""
    if isinstance(src, tuple):
        return "%s:%s" % src[:2]
    return str(src)


def _safe(text):
    """엑셀이 거부하는 제어 문자 제거"""
    if not isinstance(text, str):
        return text
    return "".join(ch if ch in "\t\n" or ord(ch) >= 32 else "·" for ch in text)


def save_excel(path, results, comm=None, title="통신 시험 보고서", source=""):
    """
    results: tab_result() 목록
    comm   : CommLog (없으면 통신 기록 시트를 만들지 않는다)
    source : 설정 파일 경로 등 보고서에 남길 출처
    """
    if openpyxl is None:
        raise ReportError("보고서를 만들려면 openpyxl 이 필요합니다: pip install openpyxl")
    wb = openpyxl.Workbook()
    thin = Side(style="thin", color="D6DCE5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=_HEAD)
    green = PatternFill("solid", fgColor=_GREEN)
    red = PatternFill("solid", fgColor=_RED)
    light = PatternFill("solid", fgColor=_LIGHT)
    wrap = Alignment(wrap_text=True, vertical="top")

    def table(ws, row, headers, widths):
        for c, (h, w) in enumerate(zip(headers, widths), 1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.font = head_font
            cell.fill = head_fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.freeze_panes = ws.cell(row=row + 1, column=1)

    def put_row(ws, row, values, fill=None, align=None):
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=c, value=_safe(v))
            cell.border = border
            if fill is not None:
                cell.fill = fill
            if align is not None:
                cell.alignment = align

    ok = overall_ok(results)

    # ---------------------------------------------------------------- 요약
    ws = wb.active
    ws.title = "요약"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16, color=_HEAD)
    info = [("판정", PASS_TEXT if ok else FAIL_TEXT), ("만든 시각", ts(time.time(), ms=False)),
            ("설정 파일", source or "-"), ("실행 PC", "%s (%s)" % (platform.node() or "-", platform.system())),
            ("연결 수", len(results)),
            ("전체 통과 / 실패", "%d / %d" % (sum(r["passed"] for r in results), sum(r["failed"] for r in results)))]
    for i, (k, v) in enumerate(info):
        ws.cell(row=3 + i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=3 + i, column=1).fill = light
        cell = ws.cell(row=3 + i, column=2, value=_safe(v))
        if k == "판정":
            cell.font = Font(bold=True, size=14, color="15803D" if ok else "B91C1C")
            cell.fill = green if ok else red
    start = 3 + len(info) + 1
    headers = ["연결", "연결 정보", "판정", "통과", "실패", "회차 / 반복", "끝난 이유", "시작", "끝", "걸린 시간"]
    table(ws, start, headers, [18, 30, 9, 8, 8, 12, 16, 24, 24, 14])
    ws.freeze_panes = None
    for i, r in enumerate(results):
        dur = (r["ended"] - r["started"]) if r["started"] and r["ended"] else None
        rep = "%d / %s" % (r["loops"], r["repeat"] or "무한")
        put_row(ws, start + 1 + i, [r["name"], r["conn"], PASS_TEXT if r["ok"] else FAIL_TEXT, r["passed"],
                                    r["failed"], rep, r["end_reason"], ts(r["started"]), ts(r["ended"]),
                                    duration_text(dur)], fill=green if r["ok"] else red)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 30

    # ---------------------------------------------------------------- 실패 목록 / 단계 결과
    step_headers = ["연결", "회차", "번호", "단계", "결과", "걸린 시간(ms)", "시각", "메시지"]
    step_widths = [16, 7, 7, 40, 8, 13, 24, 60]
    ws_fail = wb.create_sheet("실패 목록")
    table(ws_fail, 1, step_headers, step_widths)
    ws_all = wb.create_sheet("단계 결과")
    table(ws_all, 1, step_headers, step_widths)
    nf = na = 1
    for r in results:
        if not r["records"] and not r["ok"]:
            nf += 1
            put_row(ws_fail, nf, [r["name"], "", "", "(시작 못 함)", sc.FAIL, "", ts(r["started"]), r["end_reason"]],
                    fill=red, align=wrap)
        if r.get("dropped"):
            na += 1
            put_row(ws_all, na, [r["name"], "", "", "(오래된 기록 %d개 생략)" % r["dropped"], "", "", "", ""])
        for rec in r["records"]:
            values = [r["name"], rec["loop"], rec["no"], rec["desc"], rec["status"], rec["ms"], ts(rec["time"]),
                      rec["msg"]]
            failed = rec["status"] == sc.FAIL
            na += 1
            put_row(ws_all, na, values, fill=red if failed else None)
            if failed:
                nf += 1
                put_row(ws_fail, nf, values, fill=red, align=wrap)
    if nf == 1:
        ws_fail.cell(row=2, column=1, value="실패한 단계가 없습니다.")
    ws_all.auto_filter.ref = "A1:H%d" % max(na, 2)

    # ---------------------------------------------------------------- 통신 기록
    if comm is not None:
        wc = wb.create_sheet("통신 기록")
        table(wc, 1, ["시각", "연결", "구분", "HEX", "텍스트", "상대"], [24, 16, 10, 50, 40, 20])
        row = 1
        if comm.dropped:
            row += 1
            wc.cell(row=row, column=1, value="(오래된 기록 %d줄 생략 - 최근 %d줄만 보관)" % (comm.dropped, comm.limit))
        fills = {k: PatternFill("solid", fgColor=bg.lstrip("#")) for k, (_, bg) in KIND_COLORS.items()}
        kind_fonts = {k: Font(bold=True, color=fg.lstrip("#")) for k, (fg, _) in KIND_COLORS.items()}
        for t, name, kind, data, src in list(comm.rows):
            row += 1
            hx, text = _data_cells(kind, data)
            put_row(wc, row, [ts(t), name, KIND_NAMES.get(kind, kind), hx, text, _src_text(src)],
                    fill=fills.get(kind))
            if kind in kind_fonts:
                wc.cell(row=row, column=3).font = kind_fonts[kind]
        wc.auto_filter.ref = "A1:F%d" % max(row, 2)

    folder = os.path.dirname(os.path.abspath(path))
    os.makedirs(folder, exist_ok=True)
    try:
        wb.save(path)
    except OSError as e:
        raise ReportError("보고서 저장 실패 (엑셀에서 열려 있나요?): %s" % e)
    return path


def default_name(prefix="시험보고서"):
    return "%s_%s.xlsx" % (prefix, time.strftime("%Y%m%d_%H%M%S"))
